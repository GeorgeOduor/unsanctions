from datetime import datetime
from shiny import App, reactive, render, ui

DEFAULT_URL = "https://scsanctions.un.org/xml/en/al-qaida"


app_ui = ui.page_fluid(
    ui.tags.head(
        ui.tags.style("""
            body {
                background: #f3f6fb;
                color: #172b4d;
                font-family: Inter, "Segoe UI", Arial, sans-serif;
            }

            .app-header {
                background: linear-gradient(120deg, #123763, #2468ad);
                color: white;
                padding: 30px 34px;
                border-radius: 0 0 18px 18px;
                margin-bottom: 28px;
                box-shadow: 0 6px 20px rgba(20, 55, 95, .15);
            }

            .app-header h1 {
                font-size: 30px;
                font-weight: 750;
                margin: 0 0 9px;
            }

            .app-header p {
                color: #e1edff;
                margin: 0;
                font-size: 15px;
            }

            .dashboard {
                display: grid;
                grid-template-columns: minmax(300px, .9fr)
                                       minmax(360px, 1.25fr);
                gap: 24px;
                align-items: start;
                padding: 0 6px 24px;
            }

            .panel {
                background: white;
                border: 1px solid #e0e8f3;
                border-radius: 15px;
                padding: 25px;
                box-shadow: 0 5px 18px rgba(31, 60, 100, .045);
                min-width: 0;
            }

            .panel h3 {
                font-size: 19px;
                font-weight: 700;
                margin: 0 0 22px;
                color: #163b68;
            }

            .section-divider {
                border: 0;
                border-top: 1px solid #e5edf7;
                margin: 25px 0;
            }

            .form-control {
                border: 1px solid #cfdbeb;
                border-radius: 8px;
                min-height: 43px;
                box-shadow: none;
            }

            .form-control:focus {
                border-color: #337bd0;
                box-shadow: 0 0 0 3px rgba(51, 123, 208, .12);
            }

            .source-link {
                display: inline-block;
                margin: 10px 0 0;
                color: #1768c4;
                font-weight: 600;
                text-decoration: none;
            }

            .source-link:hover {
                text-decoration: underline;
            }

            .upload-note {
                color: #6b7f99;
                font-size: 13px;
                margin: -8px 0 14px;
            }

            .btn-primary {
                background: #2169bd;
                border-color: #2169bd;
                border-radius: 9px;
                font-weight: 650;
                padding: 11px 18px;
                transition: background .15s ease;
            }

            .btn-primary:hover {
                background: #174f94;
                border-color: #174f94;
            }

            .process-button {
                width: 100%;
                margin-top: 17px;
                min-height: 47px;
            }

            .notice {
                background: #eaf3ff;
                border: 1px solid #d0e3fc;
                border-radius: 12px;
                padding: 17px 20px;
                margin-bottom: 22px;
                color: #244c7c;
            }

            .notice strong {
                display: block;
                margin-bottom: 6px;
                color: #1759a1;
            }

            .stat-grid {
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 14px;
                margin: 18px 0 22px;
            }

            .stat-card {
                border-radius: 12px;
                padding: 20px;
                background: #edf5ff;
                border: 1px solid #e0edfc;
            }

            .stat-card.entity {
                background: #eaf7f4;
                border-color: #d7eee8;
            }

            .stat-label {
                color: #4f6682;
                font-size: 13px;
                font-weight: 650;
            }

            .stat-number {
                color: #174779;
                font-size: 32px;
                font-weight: 750;
                line-height: 1.3;
                margin-top: 7px;
            }

            .stat-card.entity .stat-number {
                color: #147968;
            }

            .status-box {
                background: #f7f9fc;
                border: 1px solid #e0e8f3;
                border-radius: 10px;
                padding: 17px;
                line-height: 1.6;
                overflow-wrap: anywhere;
            }

            .download-panel {
                margin-top: 22px;
                padding: 19px;
                border: 1px solid #e0e8f3;
                border-radius: 11px;
            }

            .download-panel h4 {
                color: #163b68;
                font-size: 16px;
                font-weight: 700;
                margin: 0 0 14px;
            }

            .download-panel .btn {
                border-radius: 8px;
                padding: 10px 15px;
                font-weight: 650;
            }

            .helper-text {
                color: #71839b;
                font-size: 12px;
                margin-top: 13px;
                line-height: 1.6;
            }

            .app-footer {
                margin: 0 6px 25px;
                padding: 18px 22px;
                background: #eaf1fa;
                border-radius: 11px;
                color: #536b89;
                font-size: 13px;
                line-height: 1.6;
            }

            @media (max-width: 760px) {
                .dashboard {
                    grid-template-columns: minmax(0, 1fr);
                    padding: 0;
                }

                .app-header {
                    padding: 25px 20px;
                }

                .app-header h1 {
                    font-size: 24px;
                }

                .panel {
                    padding: 20px;
                }
            }
        """)
    ),

    ui.div(
        ui.div(
            ui.h1("UN Sanctions XML Extractor"),
            ui.p(
                "Download the source XML, upload it, "
                "and export the extracted records."
            ),
            class_="app-header",
        ),

        ui.div(
            # LEFT: source and upload
            ui.div(
                ui.h3("01  ·  XML Source"),
                ui.input_text(
                    "xml_url",
                    "UN XML source URL",
                    value=DEFAULT_URL,
                    width="100%",
                ),
                ui.output_ui("source_link"),

                ui.tags.hr(class_="section-divider"),

                ui.h3("02  ·  Upload XML File"),
                ui.p(
                    "Select the XML file downloaded from the source URL.",
                    class_="upload-note",
                ),
                ui.input_file(
                    "xml_file",
                    "Choose XML file",
                    accept=[".xml", "text/xml", "application/xml"],
                    multiple=False,
                ),
                ui.input_action_button(
                    "process",
                    "Process XML",
                    class_="btn-primary process-button",
                ),
                class_="panel",
            ),

            # RIGHT: results and downloads
            ui.div(
                ui.div(
                    ui.tags.strong("Ready to process"),
                    ui.span(
                        "1. Open the XML source.  "
                        "2. Save the file.  3. Upload it here."
                    ),
                    class_="notice",
                ),

                ui.h3("Extraction Summary"),

                ui.output_ui("summary"),

                ui.tags.hr(class_="section-divider"),

                ui.h3("Processing Status"),
                ui.div(
                    ui.output_ui("status"),
                    class_="status-box",
                ),

                ui.div(
                    ui.h4("Download Excel Workbook"),
                    ui.download_button(
                        "download_excel",
                        "Download Excel workbook",
                        class_="btn-primary",
                    ),
                    ui.p(
                        "Your workbook contains two sheets: "
                        "Individuals and Entities.",
                        class_="helper-text",
                    ),
                    class_="download-panel",
                ),
                class_="panel",
            ),

            class_="dashboard",
        ),

        ui.div(
            "This tool extracts and organizes records from UN sanctions "
            "XML files into a formatted Excel workbook. Always verify "
            "screening results against the official source.",
            class_="app-footer",
        ),
    ),
    title="UN Sanctions XML Extractor",
)


def parse_xml(xml_bytes):
    import xml.etree.ElementTree as ET

    root = ET.fromstring(xml_bytes)

    def text(el, tag):
        node = el.find(tag)
        return (node.text or "").strip() if node is not None else ""

    def values(el, tag):
        parent = el.find(tag)
        if parent is None:
            return ""
        return " | ".join(
            (v.text or "").strip()
            for v in parent.findall("VALUE")
            if (v.text or "").strip()
        )

    def join_multi(items):
        return " | ".join(x for x in items if x)

    ind_rows = []
    ent_rows = []

    for ind in root.findall(".//INDIVIDUAL"):
        aliases = ind.findall("INDIVIDUAL_ALIAS")

        addresses = join_multi([
            ", ".join(filter(None, [
                text(a, "STREET"), text(a, "CITY"),
                text(a, "STATE_PROVINCE"), text(a, "COUNTRY")
            ]))
            for a in ind.findall("INDIVIDUAL_ADDRESS")
        ])

        dobs = []
        for dob in ind.findall("INDIVIDUAL_DATE_OF_BIRTH"):
            kind = text(dob, "TYPE_OF_DATE")
            if kind == "EXACT":
                dobs.append(text(dob, "DATE") or text(dob, "YEAR"))
            elif kind == "APPROXIMATELY":
                year = text(dob, "YEAR")
                if year:
                    dobs.append(f"~{year}")
            elif kind == "BETWEEN":
                start, end = text(dob, "FROM_YEAR"), text(dob, "TO_YEAR")
                if start or end:
                    dobs.append(f"{start}{end}")

        pobs = join_multi([
            ", ".join(filter(None, [
                text(p, "CITY"), text(p, "STATE_PROVINCE"),
                text(p, "COUNTRY")
            ]))
            for p in ind.findall("INDIVIDUAL_PLACE_OF_BIRTH")
        ])

        docs = join_multi([
            " ".join(filter(None, [
                text(d, "TYPE_OF_DOCUMENT"), text(d, "NUMBER"),
                text(d, "ISSUING_COUNTRY"), text(d, "DATE_OF_ISSUE")
            ]))
            for d in ind.findall("INDIVIDUAL_DOCUMENT")
        ])

        ind_rows.append([
            text(ind, "DATAID"),
            text(ind, "REFERENCE_NUMBER"),
            text(ind, "LISTED_ON"),
            text(ind, "FIRST_NAME"),
            text(ind, "SECOND_NAME"),
            text(ind, "THIRD_NAME"),
            text(ind, "FOURTH_NAME"),
            text(ind, "GENDER"),
            text(ind, "NAME_ORIGINAL_SCRIPT"),
            values(ind, "TITLE"),
            values(ind, "DESIGNATION"),
            values(ind, "NATIONALITY"),
            join_multi([text(a, "ALIAS_NAME") for a in aliases]),
            join_multi([text(a, "QUALITY") for a in aliases]),
            join_multi(dobs),
            pobs,
            addresses,
            docs,
            values(ind, "LAST_DAY_UPDATED"),
            text(ind, "COMMENTS1").replace("\n", " ")
        ])

    for ent in root.findall(".//ENTITY"):
        aliases = ent.findall("ENTITY_ALIAS")

        addresses = join_multi([
            ", ".join(filter(None, [
                text(a, "STREET"), text(a, "CITY"),
                text(a, "STATE_PROVINCE"), text(a, "ZIP_CODE"),
                text(a, "COUNTRY")
            ]))
            for a in ent.findall("ENTITY_ADDRESS")
        ])

        ent_rows.append([
            text(ent, "DATAID"),
            text(ent, "REFERENCE_NUMBER"),
            text(ent, "LISTED_ON"),
            text(ent, "FIRST_NAME") or text(ent, "NAME"),
            text(ent, "NAME_ORIGINAL_SCRIPT"),
            join_multi([text(a, "ALIAS_NAME") for a in aliases]),
            join_multi([text(a, "QUALITY") for a in aliases]),
            addresses,
            values(ent, "LAST_DAY_UPDATED"),
            text(ent, "COMMENTS1").replace("\n", " ")
        ])

    if not ind_rows and not ent_rows:
        raise ValueError(
            "No individual or entity records found in this XML file."
        )

    return ind_rows, ent_rows


def make_workbook(ind_rows, ent_rows):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ind_headers = [
        "DATAID", "REFERENCE_NUMBER", "LISTED_ON",
        "FIRST_NAME", "SECOND_NAME", "THIRD_NAME", "FOURTH_NAME",
        "GENDER", "NAME_ORIGINAL_SCRIPT", "TITLE", "DESIGNATION",
        "NATIONALITY", "ALIASES", "ALIAS_QUALITIES",
        "DATES_OF_BIRTH", "PLACES_OF_BIRTH", "ADDRESSES",
        "DOCUMENTS", "LAST_DAY_UPDATED", "COMMENTS"
    ]

    ent_headers = [
        "DATAID", "REFERENCE_NUMBER", "LISTED_ON", "NAME",
        "NAME_ORIGINAL_SCRIPT", "ALIASES", "ALIAS_QUALITIES",
        "ADDRESSES", "LAST_DAY_UPDATED", "COMMENTS"
    ]

    wb = Workbook()
    header_fill = PatternFill("solid", start_color="1F4E79")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    data_font = Font(name="Arial", size=10)

    def build_sheet(ws, title, headers, rows):
        ws.title = title
        ws.append(headers)

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center", vertical="top"
            )

        for row in rows:
            ws.append(row)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.alignment = Alignment(
                    vertical="top", wrap_text=True
                )

        for i, header in enumerate(headers, 1):
            max_len = len(header)
            for row in rows:
                max_len = max(
                    max_len, min(len(str(row[i - 1] or "")), 60)
                )
            ws.column_dimensions[get_column_letter(i)].width = min(
                max_len + 2, 62
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    build_sheet(wb.active, "Individuals", ind_headers, ind_rows)
    build_sheet(wb.create_sheet("Entities"), "Entities", ent_headers, ent_rows)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def server(input, output, session):
    result = reactive.Value(None)
    status_message = reactive.Value(
        "1. Open the XML source. 2. Save the file. 3. Upload it here."
    )
    def source_link():
            url = input.xml_url().strip()
        
            if not url.startswith(("https://", "http://")):
                return ui.p("Enter a valid HTTP or HTTPS URL.")
        
            return ui.tags.a(
                "— Open XML source in a new tab",
                href=url,
                target="_blank",
                class_="source-link",
            )
    @reactive.effect
    @reactive.event(input.process)
    def process_file():
        uploaded = input.xml_file()

        if not uploaded:
            result.set(None)
            status_message.set("Please upload an XML file first.")
            return

        result.set(None)

        try:
            with open(uploaded[0]["datapath"], "rb") as f:
                xml_bytes = f.read()

            individuals, entities = parse_xml(xml_bytes)
            workbook = make_workbook(individuals, entities)

            result.set(
                {
                    "individuals": individuals,
                    "entities": entities,
                    "workbook": workbook,
                }
            )

            status_message.set(
                f"Successfully processed {len(individuals)} individuals "
                f"and {len(entities)} entities."
            )

        except Exception as exc:
            status_message.set(f"Processing failed: {exc}")

    @render.ui
    def status():
        return ui.p(status_message.get())

    @render.ui
    def summary():
        data = result.get()
    
        individuals = len(data["individuals"]) if data else 0
        entities = len(data["entities"]) if data else 0
    
        return ui.div(
            ui.div(
                ui.div("Individuals", class_="stat-label"),
                ui.div(f"{individuals:,}", class_="stat-number"),
                class_="stat-card",
            ),
            ui.div(
                ui.div("Entities", class_="stat-label"),
                ui.div(f"{entities:,}", class_="stat-number"),
                class_="stat-card entity",
            ),
            class_="stat-grid",
        )

    @render.download(
        filename=lambda: (f"sanctions_{datetime.now().strftime('%d.%m.%Y')}.xlsx"),
        media_type=(
            "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet"
        ),
    )
    async def download_excel():
        data = result.get()
        if data is not None:
            yield data["workbook"]


app = App(app_ui, server)
